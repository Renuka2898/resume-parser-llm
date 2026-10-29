"""PDF/DOCX text extraction."""
from io import BytesIO
import fitz
from docx import Document

def extract_resume_text(filename: str, content: bytes) -> str:
    """Extract text from a PDF or DOCX, raising ValueError for invalid/empty files."""
    if filename.lower().endswith(".pdf"):
        try:
            with fitz.open(stream=content, filetype="pdf") as doc:
                text = "\n".join(page.get_text("text") for page in doc)
        except Exception as exc:
            raise ValueError("Could not read this PDF.") from exc
    elif filename.lower().endswith(".docx"):
        try:
            doc = Document(BytesIO(content))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                parts.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            text = "\n".join(parts)
        except Exception as exc:
            raise ValueError("Could not read this DOCX.") from exc
    else:
        raise ValueError("Only PDF and DOCX files are supported.")
    if not text.strip():
        raise ValueError("No readable text found. Scanned PDFs may require OCR.")
    return text.strip()
